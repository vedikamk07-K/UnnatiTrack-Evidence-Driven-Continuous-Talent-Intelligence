"""Extract plain text from an uploaded resume (PDF, DOCX or TXT)."""
from __future__ import annotations

import io

from fastapi import HTTPException

from .engine.resume import RESUME

TYPES = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".txt": "text/plain"}


def ext_of(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def extract_text(filename: str, data: bytes) -> str:
    ext = ext_of(filename)
    if ext not in TYPES:
        raise HTTPException(415, "Upload a PDF, DOCX or TXT resume.")
    if not data:
        raise HTTPException(422, "The file is empty.")
    if len(data) > RESUME["maxBytes"]:
        raise HTTPException(413, "Resume must be 5 MB or smaller.")
    try:
        if ext == ".txt":
            return data.decode("utf-8", errors="replace")
        if ext == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                parts += [c.text for row in t.rows for c in row.cells]
            return "\n".join(parts)
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    except HTTPException:
        raise
    except Exception as e:  # corrupt or password-protected files
        raise HTTPException(422, f"Could not read {ext[1:].upper()} file: {e.__class__.__name__}.") from e
